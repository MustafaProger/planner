"""Read-only inventory and reconciliation; run from the repository root.

Uses only the two business workbooks and the requested transcript. Never opens
the separate access workbook and never imports app.server (which initializes DB).
Output contains counts/schema/checksums, not individual business records.
"""
import argparse
import collections
import datetime as dt
import hashlib
import json
from pathlib import Path
import sqlite3
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
import openpyxl
from app.importer import load_excel
from app import workspace

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inventory(path):
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for sheet in formulas:
        cached = values[sheet.title]
        sheets.append(dict(
            name=sheet.title, range=sheet.calculate_dimension(), visibility=sheet.sheet_state,
            nonempty_cells=sum(c.value is not None for row in sheet for c in row),
            formulas=sum(c.data_type == 'f' for row in sheet for c in row),
            cached_errors=dict(collections.Counter(str(c.value) for row in cached for c in row if c.data_type == 'e')),
            tables=[dict(name=t.name, range=t.ref, columns=[c.name for c in t.tableColumns]) for t in sheet.tables.values()],
        ))
    with zipfile.ZipFile(path) as archive:
        pivots = []
        for name in archive.namelist():
            if name.startswith('xl/pivotTables/pivotTable') and name.endswith('.xml'):
                element = ET.fromstring(archive.read(name))
                pivots.append(dict(name=element.get('name'), cache_id=element.get('cacheId')))
        has_binary_model = any(name.startswith('xl/model/') for name in archive.namelist())
    formulas.close()
    values.close()
    return dict(file=path.name, sha256=digest(path), sheets=sheets, pivots=pivots, binary_power_pivot_model=has_binary_model)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, default=Path('/Users/mustafa/Desktop/проект'))
    parser.add_argument('--database', type=Path, default=ROOT / 'data/fleet.sqlite')
    parser.add_argument('--output', type=Path, default=Path(__file__).with_name('excel-verification.json'))
    args = parser.parse_args()
    stat = args.source / 'Статистика 17.09.2026.xlsx'
    plan = args.source / 'Планировщик ЛОГИСТИКА АА (Р4).xlsx'
    transcript = args.source / 'Новая запись 30 (mp3cut.net).txt'
    before = {p.name: digest(p) for p in (stat, plan, transcript)}
    source = load_excel(stat, plan)
    result = dict(checked_at=dt.datetime.now(dt.timezone(dt.timedelta(hours=3))).isoformat(),
                  transcript=dict(file=transcript.name, sha256=digest(transcript),
                                  lines=len(transcript.read_text().splitlines()),
                                  matches_document_copy=transcript.read_bytes() == (ROOT / 'docs/Разговор.txt').read_bytes()),
                  books=[inventory(stat), inventory(plan)],
                  source_counts=source['meta']['counts'],
                  source_categories=dict(collections.Counter(r['category'] for r in source['trips'])),
                  source_finance={f:round(sum(r.get(f) or 0 for r in source['trips']),2) for f in ('revenue','payout')},
                  source_satellite={f:round(sum(r.get(f) or 0 for r in source['satellite']),6) for f in ('km','hours','engine_hours','movement_hours','stop_hours','idle_hours')})
    connection = sqlite3.connect(args.database.resolve().as_uri() + '?mode=ro', uri=True)
    connection.row_factory = sqlite3.Row
    connection.execute('begin')
    snapshot = json.loads(connection.execute('select body from snapshot where id=1').fetchone()[0])
    current = workspace.dataset(connection)
    preserved = {}
    for kind in ('plans','statuses','satellite'):
        old = {r['source'] for r in source[kind]}
        now = {r['source'] for r in current[kind]}
        preserved[kind] = dict(source_rows=len(source[kind]), missing_source_references=len(old-now))
    same_snapshot = {kind:source[kind] == snapshot[kind] for kind in ('trips','plans','statuses','routes')}
    snapshot_satellite = {r['source']:r for r in snapshot['satellite']}
    common_satellite_fields_unchanged = all(
        all(row.get(key) == value for key,value in snapshot_satellite[row['source']].items())
        for row in source['satellite']
    )
    result['working_database'] = dict(
        integrity_check=connection.execute('pragma integrity_check').fetchone()[0],
        counts=current['meta']['counts'],
        periods={kind:dict(sorted(collections.Counter(r['date'][:7] for r in current[kind]).items())) for kind in ('trips','plans','statuses','satellite')},
        catalogs={kind:len(rows) for kind,rows in current['directories'].items()},
        empty_business_tables={kind:len(current[kind]) for kind in ('coverage','card_bindings','expenses')},
        trip_origins=dict(collections.Counter(r.get('origin') for r in current['trips'])),
        missing_trip_metrics={field:sum(r.get(field) is None for r in current['trips']) for field in ('km','hours','revenue','payout')},
        plan_states=dict(collections.Counter(r['state'] for r in current['plans'])),
        plans_flagged_duplicate=sum(bool(r.get('duplicate_review')) for r in current['plans']),
        source_references_preserved=preserved,
        snapshot_arrays_match_source=same_snapshot,
        original_satellite_fields_match_source=common_satellite_fields_unchanged,
        source_hashes_match_original_snapshot=source['meta']['files'] == snapshot['meta']['files'],
    )
    connection.rollback()
    connection.close()
    result['source_files_unchanged'] = before == {p.name:digest(p) for p in (stat, plan, transcript)}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(dict(output=str(args.output), books=len(result['books']),
                          sheets=sum(len(b['sheets']) for b in result['books']),
                          tables=sum(len(s['tables']) for b in result['books'] for s in b['sheets']),
                          pivots=sum(len(b['pivots']) for b in result['books']),
                          source_files_unchanged=result['source_files_unchanged']), ensure_ascii=False))


if __name__ == '__main__':
    main()
